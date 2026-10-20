"""把 materials/templates/*.csv 合成一个 Excel 工作簿，方便发给甲方填写。

用法：pip install openpyxl && python3 tools/build_materials_xlsx.py
输出：materials/红河同路-资料收集表.xlsx
"""
import csv
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TPL = os.path.join(ROOT, "materials", "templates")
OUT = os.path.join(ROOT, "materials", "红河同路-资料收集表.xlsx")

# 与 materials/README.md 第 1 节顺序一致
ORDER = [
    "站点信息", "内容卡", "题目", "石榴籽与非遗卡", "音频清单", "县情数据", "问答库",
    "礼品与档位", "柜台与店员", "合作社", "商品与来源档案", "素材授权登记", "服务点与电话", "审定人员",
]

HEAD = PatternFill("solid", fgColor="1F5F5B")
NOTE = PatternFill("solid", fgColor="F3EBDD")
SAMPLE = PatternFill("solid", fgColor="FAFAF7")


def read_readme_index():
    """从 README 第 1 节表格取「提供单位 / 需要时间 / 不给会卡住」，写进目录页。"""
    path = os.path.join(ROOT, "materials", "README.md")
    rows = []
    with open(path, encoding="utf-8") as f:
        for line in f:
            if re.match(r"^\| \d+ \|", line):
                cells = [c.strip() for c in line.strip().strip("|").split("|")]
                rows.append(cells)
    return rows


def main():
    wb = Workbook()
    idx = wb.active
    idx.title = "目录"
    idx.append(["#", "资料", "模板", "提供单位", "需要时间", "不给会卡住"])
    for r in read_readme_index():
        idx.append([re.sub(r"[`\[\]]|\(.*?\)", "", c) for c in r])
    for c in idx[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
    for col, w in zip("ABCDEF", (5, 36, 32, 28, 26, 26)):
        idx.column_dimensions[col].width = w
    for row in idx.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    for name in ORDER:
        path = os.path.join(TPL, name + ".csv")
        with open(path, encoding="utf-8-sig") as f:
            rows = list(csv.reader(f))
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
        ws.freeze_panes = "A3"
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = HEAD
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for c in ws[2]:
            c.font = Font(color="7A5A2A", size=9)
            c.fill = NOTE
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for row in ws.iter_rows(min_row=3, max_row=len(rows)):
            for c in row:
                c.fill = SAMPLE
                c.font = Font(color="888888", italic=True)
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, h in enumerate(rows[0], start=1):
            longest = max(len(r[i - 1]) if i - 1 < len(r) else 0 for r in rows)
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(40, longest * 1.6 if longest < 12 else 30))
        ws.row_dimensions[2].height = 60

    wb.save(OUT)
    print("已生成", os.path.relpath(OUT, ROOT))


if __name__ == "__main__":
    main()
